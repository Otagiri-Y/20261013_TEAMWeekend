# extract.py の先頭に書く形
import io
from pathlib import Path
from docx import Document
from pypdf import PdfReader

def extract_text(file_name: str, data: bytes) -> str:
    ext = Path(file_name).suffix.lower()

    if ext == ".txt":
        for enc in ("utf-8-sig", "cp932"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        raise ValueError("文字コードを判別できません（UTF-8 か Shift_JIS で保存してください）")

    if ext == ".docx":
        try:
            doc = Document(io.BytesIO(data))
        except Exception as e:
            raise ValueError(f"docx を開けません: {e}")
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(lines)

    if ext == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as e:
            raise ValueError(f"PDF を読み込めません: {e}")
        if not text.strip():
            raise ValueError("PDF から文字を取り出せません（スキャン画像の可能性があります）")
        return text

    raise ValueError(f"未対応の形式です: {ext}（対応: txt / docx / pdf）")